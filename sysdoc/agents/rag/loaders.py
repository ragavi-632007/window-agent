# RAG document loaders
"""Type-aware loaders for local files and folders."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable, List, Optional

from langchain_core.documents import Document

# Extensions we attempt to load
_TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".json", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".env.example",
}
_CODE_EXT = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".scala", ".sh",
    ".ps1", ".bat", ".sql", ".html", ".css", ".scss", ".vue", ".svelte",
}
_DOC_EXT = {".pdf", ".docx"}
_SKIP_DIR_NAMES = {
    "node_modules", ".git", "__pycache__", ".venv", "venv", "dist", "build",
    ".next", ".cache", "target", "bin", "obj", "$Recycle.Bin",
}

_MAX_FILE_BYTES = 2 * 1024 * 1024  # 2 MB per file hard cap for MVP


def _read_text(path: Path) -> Optional[str]:
    try:
        size = path.stat().st_size
        if size > _MAX_FILE_BYTES or size == 0:
            return None
        data = path.read_bytes()
        # skip obvious binary
        if b"\x00" in data[:4096]:
            return None
        return data.decode("utf-8", errors="replace")
    except Exception:
        return None


def load_text_file(path: str | Path) -> List[Document]:
    p = Path(path)
    text = _read_text(p)
    if not text:
        return []
    return [Document(page_content=text, metadata={"source": str(p), "file_type": p.suffix.lower() or "text"})]


def load_pdf(path: str | Path) -> List[Document]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    p = Path(path)
    docs: List[Document] = []
    try:
        reader = PdfReader(str(p))
        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={"source": str(p), "page": i + 1, "file_type": "pdf"},
                ))
    except Exception:
        return []
    return docs


def load_docx(path: str | Path) -> List[Document]:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return []
    p = Path(path)
    try:
        doc = DocxDocument(str(p))
        parts = [para.text for para in doc.paragraphs if para.text and para.text.strip()]
        text = "\n".join(parts)
        if not text.strip():
            return []
        return [Document(page_content=text, metadata={"source": str(p), "file_type": "docx"})]
    except Exception:
        return []


def load_single(path: str | Path) -> List[Document]:
    p = Path(path)
    if not p.is_file():
        return []
    ext = p.suffix.lower()
    if ext == ".pdf":
        return load_pdf(p)
    if ext == ".docx":
        return load_docx(p)
    if ext in _TEXT_EXT or ext in _CODE_EXT or ext == "":
        return load_text_file(p)
    # try text anyway for unknown text-like
    return load_text_file(p)


def iter_folder_files(
    root: str | Path,
    max_files: int = 80,
    max_depth: int = 6,
) -> Iterable[Path]:
    root = Path(root)
    count = 0
    for dirpath, dirnames, filenames in os.walk(root):
        depth = Path(dirpath).relative_to(root).parts
        if len(depth) > max_depth:
            dirnames.clear()
            continue
        # prune noisy dirs
        dirnames[:] = [d for d in dirnames if d not in _SKIP_DIR_NAMES and not d.startswith(".")]
        for name in filenames:
            if count >= max_files:
                return
            p = Path(dirpath) / name
            ext = p.suffix.lower()
            if ext in _TEXT_EXT or ext in _CODE_EXT or ext in _DOC_EXT:
                yield p
                count += 1


def load_path(path: str | Path, max_files: int = 80) -> List[Document]:
    """Load a file or walk a folder into Documents."""
    p = Path(path)
    if p.is_file():
        return load_single(p)
    if p.is_dir():
        docs: List[Document] = []
        for fp in iter_folder_files(p, max_files=max_files):
            docs.extend(load_single(fp))
        return docs
    return []


def estimate_scope(path: str | Path, max_files: int = 80) -> dict:
    """Cheap estimate for permission gate."""
    p = Path(path)
    if p.is_file():
        try:
            mb = p.stat().st_size / (1024 * 1024)
        except Exception:
            mb = 0.0
        return {"files": 1, "mb": round(mb, 2), "paths": [str(p)]}
    files = list(iter_folder_files(p, max_files=max_files))
    total = 0
    for f in files:
        try:
            total += f.stat().st_size
        except Exception:
            pass
    return {
        "files": len(files),
        "mb": round(total / (1024 * 1024), 2),
        "paths": [str(f) for f in files[:20]],
    }